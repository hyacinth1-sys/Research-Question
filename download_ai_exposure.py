#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""下载并计算2021—2024年美国县级AI暴露度。

Python 3.10+；必需 openpyxl，pandas 可选（用于输出Stata文件）。
安装：python -m pip install openpyxl pandas
运行：python download_ai_exposure.py
离线重算：python download_ai_exposure.py --offline

结果默认保存在代码旁的 AI_exposure_2021_2024 文件夹。
口径：50州+DC，QCEW私营部门年度平均就业，披露且评分匹配的行业归一化权重。
处理NAICS2017/2022跨分类；共同组评分使用固定的2021全国就业权重。
分数代表技术暴露潜力，年度变化来自就业结构；不是实际AI采用率。
"""
from __future__ import annotations
import argparse
import csv
from datetime import datetime, timezone
import hashlib
from http.client import HTTPException
import io
import json
import logging
import math
import os
from pathlib import Path
import re
import shutil
import socket
import statistics
import subprocess
import sys
import time
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
import zipfile

LOGGER = logging.getLogger('county_ai_exposure')
BLOCK_SIZE = 1024 * 1024

def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()

def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(BLOCK_SIZE), b""):
            digest.update(block)
    return digest.hexdigest()

def write_json(path: Path, value: dict) -> None:
    """先写临时文件，成功后替换；中断不留下假完整JSON。"""
    temporary = path.with_name(path.name + ".part")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    os.replace(temporary, path)

def request(url: str, method: str = "GET") -> Request:
    return Request(
        url, method=method,
        headers={"User-Agent": "County-AI-Exposure-Downloader/1.0", "Accept-Encoding": "identity"},
    )

class CurlDownloadError(OSError):
    def __init__(self, returncode: int, status: int, message: str):
        super().__init__(f"curl退出码{returncode}，HTTP{status}：{message}")
        self.returncode, self.status = returncode, status

def curl_request(url: str, timeout: float, target: Path | None = None, year: int | None = None) -> dict:
    """通过同一官方URL下载；不调用shell，不改用其他数据源。"""
    executable = shutil.which("curl.exe" if os.name == "nt" else "curl")
    if executable is None:
        raise ValueError("找不到curl下载工具。请安装curl或在Python能直接访问数据源的网络重试。")
    marker = "\nAI_TRANSFER_INFO\n"
    command = [
        executable, "--location", "--silent", "--show-error", "--fail",
        "--proto", "=https", "--proto-redir", "=https",
        "--connect-timeout", str(min(timeout, 60)),
        "--speed-limit", "1", "--speed-time", str(max(1, math.ceil(timeout))),
        "--dump-header", "-", "--output", str(target) if target else os.devnull,
        "--write-out", marker + "%{http_code}\n%{url_effective}\n%{size_download}\n",
    ]
    if target is None:
        command.append("--head")
    elif target.exists():
        target.unlink()  # 仅清理本次未完成的.part文件。
    command.append(url)
    process = subprocess.Popen(
        command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
        encoding="utf-8", errors="replace", shell=False,
        creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
    )
    try:
        while True:
            try:
                stdout, stderr = process.communicate(timeout=10)
                break
            except subprocess.TimeoutExpired:
                if target is not None and target.exists():
                    LOGGER.info("%s已下载%s MiB（curl）。", year, f"{target.stat().st_size / BLOCK_SIZE:,.1f}")
    except BaseException:
        process.terminate()
        try:
            process.communicate(timeout=5)
        except subprocess.TimeoutExpired:
            process.kill()
            process.communicate()
        raise
    headers_text, separator, information = stdout.rpartition(marker)
    parts = information.strip().splitlines() if separator else []
    status = int(parts[0]) if parts and parts[0].isdigit() else 0
    if process.returncode != 0:
        raise CurlDownloadError(process.returncode, status, stderr.strip())
    if len(parts) != 3 or status != 200:
        raise ValueError(f"curl未返回完整HTTP200响应：{status}。")
    # 多跳重定向、代理CONNECT或103响应，只使用最终HTTP响应头。
    headers: dict[str, str] = {}
    header_status = 0
    for line in headers_text.replace("\r\n", "\n").splitlines():
        if line.startswith("HTTP/"):
            headers = {}
            words = line.split()
            header_status = int(words[1]) if len(words) > 1 and words[1].isdigit() else 0
        elif ":" in line and header_status:
            key, value = line.split(":", 1)
            headers[key.lower().strip()] = value.strip()
    if header_status != status:
        raise ValueError("curl响应头状态与最终HTTP状态不一致。")
    return {
        "resolved_url": parts[1], "headers": headers,
        "reported_bytes": int(float(parts[2])), "download_transport": "curl",
    }

def download_using_urllib(url: str, temporary: Path, timeout: float, year: int) -> dict:
    digest = hashlib.sha256()
    received = 0
    last_report = time.monotonic()
    with urlopen(request(url), timeout=timeout) as response, temporary.open("wb") as handle:
        if response.status != 200:
            raise ValueError(f"需要完整文件的HTTP200响应，得到{response.status}。")
        headers = {key.lower(): value for key, value in response.headers.items()}
        expected = headers.get("content-length")
        for block in iter(lambda: response.read(BLOCK_SIZE), b""):
            handle.write(block)
            digest.update(block)
            received += len(block)
            if time.monotonic() - last_report >= 10:
                progress = f"/{int(expected) / BLOCK_SIZE:,.1f} MiB" if expected is not None else " MiB"
                LOGGER.info("%s已下载%s%s", year, f"{received / BLOCK_SIZE:,.1f}", progress)
                last_report = time.monotonic()
        return {
            "resolved_url": response.geturl(), "headers": headers,
            "bytes": received, "sha256": digest.hexdigest(), "download_transport": "urllib",
        }

def fetch(url: str, target: Path, timeout: float, retries: int, refresh: bool) -> dict:
    """缓存通过来源与SHA核验才复用；每次完整下载后才发布正式原始文件。"""
    receipt = target.with_name(target.name + ".source.json")
    temporary = target.with_name(target.name + ".part")
    if target.exists() and not refresh:
        if not receipt.exists():
            raise ValueError(f"缓存缺少来源记录：{target}；请加--refresh。")
        metadata = json.loads(receipt.read_text(encoding="utf-8"))
        if metadata.get("source_url") != url or metadata.get("sha256") != sha256_file(target):
            raise ValueError(f"缓存与来源不一致：{target}；请加--refresh。")
        LOGGER.info("复用校验通过的文件：%s", target.name)
        return metadata
    for attempt in range(1, retries + 1):
        try:
            LOGGER.info("下载%s（%s/%s）", target.name, attempt, retries)
            try:
                result = download_using_urllib(url, temporary, timeout, target.name)
            except HTTPError as exc:
                if exc.code != 403:
                    raise
                LOGGER.info("Python请求被拒绝，改用curl访问相同URL。")
                result = curl_request(url, timeout, temporary, target.name)
                result["bytes"] = temporary.stat().st_size
                result["sha256"] = sha256_file(temporary)
                if result["reported_bytes"] != result["bytes"]:
                    raise HTTPException("curl报告字节数与文件大小不一致。")
            expected = result["headers"].get("content-length")
            if expected is not None and int(expected) != result["bytes"]:
                raise HTTPException("下载长度不完整。")
            with temporary.open("rb") as stream:
                beginning = stream.read(4096).lstrip().lower()
            if not beginning or beginning.startswith((b"<!doctype", b"<html")):
                raise ValueError("下载结果为空或为错误网页。")
            metadata = {"source_url": url, "resolved_url": result["resolved_url"],
                        "downloaded_at_utc": utc_now(), "bytes": result["bytes"],
                        "sha256": result["sha256"], "transport": result["download_transport"],
                        "etag": result["headers"].get("etag"),
                        "last_modified": result["headers"].get("last-modified")}
            os.replace(temporary, target)
            write_json(receipt, metadata)
            return metadata
        except HTTPError as exc:
            if exc.code not in {408, 429, 500, 502, 503, 504} or attempt == retries:
                raise
        except CurlDownloadError as exc:
            retryable = exc.status in {408, 429, 500, 502, 503, 504} or (
                exc.status < 400 and exc.returncode in {5, 6, 7, 18, 28, 35, 52, 55, 56, 92})
            if not retryable or attempt == retries:
                raise
        except (URLError, TimeoutError, socket.timeout, HTTPException):
            if attempt == retries:
                raise
        LOGGER.warning("暂时的网络错误，稍后从头重试。")
        time.sleep(min(2 ** attempt, 30))
    raise RuntimeError("未完成下载。")

"""Verified Felten AIIE and Census NAICS 2017/2022 common-classification readers.

This module only reads downloaded workbooks; no downloads or employment allocation.
Use load_classification(raw_dir). Values remain in the authors' published units.
"""

import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path

import openpyxl


def normalized_title(value):
    return " ".join(str(value).split()).casefold()


# Exact published OEWS aggregation descriptions. These assign one published score
# to its stated member industries. They never replicate any employment row.
_COMPOSITE_TITLES = {
    "Chemical Manufacturing (3255 and 3256 only)": ["3255", "3256"],
    "Chemical Manufacturing (3251, 3252, 3253, and 3259 only)": ["3251", "3252", "3253", "3259"],
    "Nonmetallic Mineral Product Manufacturing": ["3271", "3272", "3273", "3274", "3279"],
    "Fabricated Metal Product Manufacturing (3323 and 3324 only)": ["3323", "3324"],
    "Fabricated Metal Product Manufacturing (3321, 3322, 3325, 3326, and 3329 only)": ["3321", "3322", "3325", "3326", "3329"],
    "Machinery Manufacturing (3331, 3332, 3334, and 3339 only)": ["3331", "3332", "3334", "3339"],
    "Furniture and Related Product Manufacturing (3371 and 3372 only)": ["3371", "3372"],
    "Merchant Wholesalers, Durable Goods (4232, 4233, 4235, 4236, 4237, and 4239 only)": ["4232", "4233", "4235", "4236", "4237", "4239"],
    "Merchant Wholesalers, Nondurable Goods (4244 and 4248 only)": ["4244", "4248"],
    "Merchant Wholesalers, Nondurable Goods (4241, 4247, and 4249 only)": ["4241", "4247", "4249"],
    "Merchant Wholesalers, Nondurable Goods (4242 and 4246 only)": ["4242", "4246"],
    "Food and Beverage Stores (4451 and 4452 only)": ["4451", "4452"],
    "General Merchandise Stores": ["4522", "4523"],
    "Miscellaneous Store Retailers (4532 and 4533 only)": ["4532", "4533"],
    "Truck Transportation": ["4841", "4842"],
    "Telecommunications": ["5173", "5174", "5179"],
    "Credit Intermediation and Related Activities (5221 And 5223 only)": ["5221", "5223"],
    "Securities, Commodity Contracts, and Other Financial Investments and Related Activities": ["5231", "5232", "5239"],
    "Real Estate": ["5311", "5312", "5313"],
    "Rental and Leasing Services (5322, 5323, and 5324 only)": ["5322", "5323", "5324"],
}
COMPOSITE_MEMBERS = {normalized_title(k): v for k, v in _COMPOSITE_TITLES.items()}


def _code_string(value):
    if isinstance(value, bool):
        raise ValueError("Boolean is not an industry code")
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value != int(value):
            raise ValueError(f"Invalid numeric industry code: {value!r}")
        return str(int(value))
    return str(value).strip()


def _check(condition, message):
    if not condition:
        raise ValueError(message)


def load_common_groups(path):
    """Return unique old/new mappings from complete official six-digit crosswalk."""
    path = Path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = "2017 to 2022 NAICS U.S."
        _check(sheet_name in workbook.sheetnames, "Official concordance sheet missing")
        rows = list(workbook[sheet_name].values)
    finally:
        workbook.close()
    _check(len(rows) >= 4, "Concordance has no data")
    _check(rows[0][0] == "2017 NAICS U.S. Matched to 2022 NAICS U.S. (Full Concordance)", "Unexpected concordance title")
    expected = ["2017 NAICS Code", "2017 NAICS Title\n(and specific piece of the 2017 industry that is contained in the 2022 industry)", "2022 NAICS Code", "2022 NAICS Title"]
    _check(list(rows[2][:4]) == expected, "Unexpected concordance headers")
    parent = {}

    def find(node):
        parent.setdefault(node, node)
        if parent[node] != node:
            parent[node] = find(parent[node])
        return parent[node]

    edges = set()
    full_rows = []
    for row_number, row in enumerate(rows[3:], 4):
        if all(v is None for v in row):
            continue
        old, new = _code_string(row[0]), _code_string(row[2])
        _check(old.isdigit() and len(old) == 6 and new.isdigit() and len(new) == 6,
               f"Unexpected concordance code at row {row_number}")
        _check(isinstance(row[1], str) and isinstance(row[3], str), f"Missing concordance title at row {row_number}")
        old4, new4 = old[:4], new[:4]
        edges.add((old4, new4))
        old_node, new_node = "2017:" + old4, "2022:" + new4
        parent[find(old_node)] = find(new_node)
        full_rows.append({"old6": old, "old_title": row[1], "new6": new, "new_title": row[3]})
    components = defaultdict(lambda: {"old": [], "new": []})
    for node in parent:
        components[find(node)]["old" if node.startswith("2017:") else "new"].append(node[5:])
    ordered = sorted(components.values(), key=lambda g: tuple(sorted(g["old"])))
    old_to_group, new_to_group, group_members = {}, {}, {}
    for index, component in enumerate(ordered, 1):
        group = f"G{index:03d}"
        old_members, new_members = sorted(component["old"]), sorted(component["new"])
        _check(old_members and new_members, f"One-sided component {group}")
        for code in old_members:
            _check(code not in old_to_group, f"Duplicate old-code mapping {code}")
            old_to_group[code] = group
        for code in new_members:
            _check(code not in new_to_group, f"Duplicate new-code mapping {code}")
            new_to_group[code] = group
        group_members[group] = {"old": old_members, "new": new_members,
                                "changed": old_members != new_members or len(old_members) > 1}
    _check(len(full_rows) == 1150 and len(edges) == 355, "Unexpected full-concordance dimensions")
    _check(len(old_to_group) == 311 and len(new_to_group) == 308 and len(group_members) == 283,
           "Unexpected four-digit common classification dimensions")
    _check(all(old_to_group[old] == new_to_group[new] for old, new in edges), "Crosswalk edge crosses common groups")
    return old_to_group, new_to_group, group_members, full_rows


def load_score_table(path, sheet_name, expected_header, old_codes, measure):
    """Parse an authored table; preserve each source title and mapping rule."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        _check(sheet_name in workbook.sheetnames, f"Missing score sheet {sheet_name}")
        rows = list(workbook[sheet_name].values)
    finally:
        workbook.close()
    _check(list(rows[0]) == list(expected_header), f"Unexpected {sheet_name} headers")
    old_codes = set(old_codes)
    values, audit_rows = {}, []
    for row_number, row in enumerate(rows[1:], 2):
        if all(v is None for v in row):
            continue
        _check(len(row) == 3, f"Unexpected score row width {sheet_name}:{row_number}")
        raw_code, title, value = row
        code = _code_string(raw_code)
        prefix = code[:4]
        _check(isinstance(title, str) and title.strip(), f"Missing score title {sheet_name}:{row_number}")
        _check(isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value),
               f"Invalid score {sheet_name}:{row_number}")
        _check(prefix.isdigit() and len(prefix) == 4, f"Invalid score code {code}")
        if measure == "broad":
            _check(code.isdigit() and len(code) == 4, f"Invalid broad code {code}")
        else:
            _check(len(code) == 6, f"Invalid language-model code {code}")
        title_key = normalized_title(title)
        if prefix in {"9991", "9992", "9993"}:
            members, rule = [], "government_OEWS_designation_excluded_private_scope"
        elif title_key in COMPOSITE_MEMBERS:
            members, rule = list(COMPOSITE_MEMBERS[title_key]), "published_OEWS_group_title"
        elif prefix in old_codes:
            members, rule = [prefix], "direct_four_digit_NAICS"
        else:
            raise ValueError(f"Unknown OEWS composite score: {code!r}, {title!r}")
        for member in members:
            _check(member in old_codes, f"Score mapped outside NAICS 2017: {member}")
            _check(member not in values, f"Multiple score assignments to {member} in {sheet_name}")
            values[member] = {"score": float(value), "source_code": code,
                              "source_title": title, "source_row": row_number, "mapping_rule": rule}
        audit_rows.append({"measure": measure, "sheet": sheet_name, "source_row": row_number,
                           "source_code": code, "source_title": title, "score": float(value),
                           "old_four_digit_members": sorted(members), "mapping_rule": rule})
    _check(len(audit_rows) == 250 and len(values) == 286, f"Unexpected {measure} score coverage")
    _check(sum(not row["old_four_digit_members"] for row in audit_rows) == 3, "Unexpected excluded score rows")
    return values, audit_rows


def load_classification(raw_dir):
    """Return score_by_old, old/new_to_group, group_members and full mapping audit."""
    raw_dir = Path(raw_dir)
    concordance = raw_dir / "2017_to_2022_NAICS.xlsx"
    broad_path = raw_dir / "AIOE_DataAppendix.xlsx"
    lm_path = raw_dir / "Language_Modeling_AIOE_AIIE.xlsx"
    old_to_group, new_to_group, group_members, crosswalk_rows = load_common_groups(concordance)
    broad, broad_rows = load_score_table(broad_path, "Appendix B", ["NAICS", "Industry Title", "AIIE"], old_to_group, "broad")
    lm, lm_rows = load_score_table(lm_path, "LM AIIE", ["NAICS", "NAICS Description", "Language Modeling AIIE"], old_to_group, "lm")
    _check(set(broad) == set(lm), "Broad and language-model score coverage differ")
    score_by_old = {code: {"broad": broad[code]["score"], "lm": lm[code]["score"],
                           "broad_source": broad[code], "lm_source": lm[code]}
                    for code in sorted(broad)}
    mixed_coverage_groups = []
    for group, members in group_members.items():
        coverage = [code in score_by_old for code in members["old"]]
        if any(coverage) and not all(coverage):
            mixed_coverage_groups.append(group)
        members["all_old_members_scored"] = all(coverage)
    _check(not mixed_coverage_groups, "A common group mixes scored and unscored old industries")
    sources = [{"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
               for path in [concordance, broad_path, lm_path]]
    return {
        "score_by_old": score_by_old,
        "old_to_group": old_to_group,
        "new_to_group": new_to_group,
        "group_members": group_members,
        "score_rows": broad_rows + lm_rows,
        "crosswalk_rows": crosswalk_rows,
        "sources": sources,
        "summary": {"crosswalk_six_digit_rows": len(crosswalk_rows), "old_four_digit_codes": len(old_to_group),
                    "new_four_digit_codes": len(new_to_group), "common_groups": len(group_members),
                    "scored_old_codes": len(score_by_old),
                    "scored_common_groups": sum(v["all_old_members_scored"] for v in group_members.values()),
                    "changed_groups": sum(v["changed"] for v in group_members.values()),
                    "unscored_old_codes": sorted(set(old_to_group) - set(score_by_old)),
                    "mixed_coverage_groups": mixed_coverage_groups},
    }




SOURCES = {
    'AIOE_DataAppendix.xlsx': 'https://raw.githubusercontent.com/AIOE-Data/AIOE/main/AIOE_DataAppendix.xlsx',
    'Language_Modeling_AIOE_AIIE.xlsx': 'https://raw.githubusercontent.com/AIOE-Data/AIOE/main/Language%20Modeling%20AIOE%20and%20AIIE.xlsx',
    '2017_to_2022_NAICS.xlsx': 'https://www.census.gov/naics/concordances/2017_to_2022_NAICS.xlsx',
    'area_titles.csv': 'https://data.bls.gov/cew/doc/titles/area/area_titles.csv',
    **{f'{y}_annual_singlefile.zip': f'https://data.bls.gov/cew/data/files/{y}/csv/{y}_annual_singlefile.zip' for y in range(2021, 2025)},
}
STATE_FIPS = set('01 02 04 05 06 08 09 10 11 12 13 15 16 17 18 19 20 21 22 23 24 25 26 27 28 29 30 31 32 33 34 35 36 37 38 39 40 41 42 44 45 46 47 48 49 50 51 53 54 55 56'.split())


def write_csv(path, fields, records):
    temporary = path.with_name(path.name + '.part')
    with temporary.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction='raise')
        writer.writeheader()
        writer.writerows(records)
    os.replace(temporary, path)


def employment(value, disclosure):
    # QCEW 未披露行中的0是占位符，不能当成零就业。
    if disclosure == 'N':
        return None
    if disclosure != '':
        raise ValueError(f'未知披露标记：{disclosure!r}')
    number = int(value)
    if number < 0:
        raise ValueError('就业人数为负。')
    return number


def county_area(code):
    return bool(re.fullmatch(r'\d{5}', code) and code[:2] in STATE_FIPS
                and code[2:] not in {'000', '999'})


def extract_qcew(path, year, processed, source_hash, refresh=False):
    """只读取私营部门的县总数、县四位行业和全国四位行业；避免层级重复。"""
    cache = processed / f'qcew_private_{year}.json'
    if cache.exists() and not refresh:
        value = json.loads(cache.read_text(encoding='utf-8'))
        receipt = cache.with_suffix('.source.json')
        meta = json.loads(receipt.read_text(encoding='utf-8')) if receipt.exists() else {}
        if meta.get('source_sha256') != source_hash or meta.get('sha256') != sha256_file(cache):
            raise ValueError(f'派生缓存校验失败：{cache}；请使用--refresh。')
        if meta.get('extractor_version') == 2:
            LOGGER.info('复用就业提取结果：%s', year)
            return value
        LOGGER.info('更新%s年的就业提取版本。', year)
    county_totals, county_industries, county_sectors, national = {}, {}, {}, {}
    last_report = time.monotonic()
    with zipfile.ZipFile(path) as archive:
        members = [n for n in archive.namelist() if n.endswith('.csv')]
        if members != [f'{year}.annual.singlefile.csv']:
            raise ValueError(f'年度ZIP内容与预期不符：{members}')
        with io.TextIOWrapper(archive.open(members[0]), encoding='utf-8-sig', newline='') as stream:
            reader = csv.reader(stream, strict=True)
            headers = next(reader)
            required = ['area_fips', 'own_code', 'industry_code', 'agglvl_code', 'size_code',
                        'year', 'qtr', 'disclosure_code', 'annual_avg_emplvl']
            if len(set(headers)) != len(headers) or not set(required) <= set(headers):
                raise ValueError('QCEW字段缺失或重复。')
            ix = {k: headers.index(k) for k in required}
            for count, row in enumerate(reader, 1):
                if len(row) != len(headers):
                    raise ValueError(f'QCEW第{count + 1}行列数异常。')
                if row[ix['own_code']] != '5' or row[ix['size_code']] != '0':
                    continue
                agg = row[ix['agglvl_code']]
                if agg not in {'71', '74', '76', '16'}:
                    continue
                if row[ix['year']] != str(year) or row[ix['qtr']] != 'A':
                    raise ValueError('年或年度频率与文件名不符。')
                area, industry = row[ix['area_fips']], row[ix['industry_code']]
                e = employment(row[ix['annual_avg_emplvl']], row[ix['disclosure_code']])
                if agg == '16' and area == 'US000':
                    if not re.fullmatch(r'\d{4}', industry) or industry in national:
                        raise ValueError('全国四位行业重复或代码异常。')
                    national[industry] = e
                elif county_area(area):
                    if agg == '71':
                        if industry != '10' or area in county_totals:
                            raise ValueError('县总就业重复或行业代码异常。')
                        county_totals[area] = e
                    elif agg == '74':
                        rows = county_sectors.setdefault(area, {})
                        if not re.fullmatch(r'\d{2}|31-33|44-45|48-49', industry) or industry in rows:
                            raise ValueError('县行业大类重复或代码异常。')
                        rows[industry] = e
                    elif agg == '76':
                        rows = county_industries.setdefault(area, {})
                        if not re.fullmatch(r'\d{4}', industry) or industry in rows:
                            raise ValueError('县四位行业重复或代码异常。')
                        rows[industry] = e
                if time.monotonic() - last_report >= 15:
                    LOGGER.info('正在读取%s年就业文件，已读取%s行。', year, f'{count:,}')
                    last_report = time.monotonic()
            # 读到EOF，zipfile自动核验该CSV的CRC。
    if not county_totals or not national or set(county_industries) - set(county_totals):
        raise ValueError('县总就业/全国行业数据不完整，或县明细没有对应总数。')
    result = {'year': year, 'county_totals': county_totals,
              'county_industries': county_industries, 'county_sectors': county_sectors, 'national': national,
              'source_rows': count}
    write_json(cache, result)
    write_json(cache.with_suffix('.source.json'), {'source_sha256': source_hash,
               'sha256': sha256_file(cache), 'extractor_version': 2, 'created_at_utc': utc_now()})
    LOGGER.info('%s年提取完成：%s个县级单位。', year, len(county_totals))
    return result


def read_area_titles(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        if not {'area_fips', 'area_title'} <= set(reader.fieldnames or []):
            raise ValueError('地区名称表字段异常。')
        result = {}
        for row in reader:
            if row['area_fips'] in result and result[row['area_fips']] != row['area_title']:
                raise ValueError('同一地区代码有冲突名称。')
            result[row['area_fips']] = row['area_title']
    return result


def fixed_group_scores(classification, baseline_national):
    """把跨版本连接组件作为共同产业，2021全国私营就业权重永远固定。"""
    fixed, audit = {}, []
    for group, members in classification['group_members'].items():
        old = members['old']
        reason = ''
        if any(k not in classification['score_by_old'] for k in old):
            reason = 'unscored_old_industry'
        elif len(old) == 1:
            fixed[group] = {metric: classification['score_by_old'][old[0]][metric]
                            for metric in ['broad', 'lm']}
        elif any(baseline_national.get(k) is None for k in old):
            reason = 'missing_2021_national_weight'
        else:
            denominator = sum(baseline_national[k] for k in old)
            if denominator <= 0:
                reason = 'zero_2021_national_weight'
            else:
                fixed[group] = {metric: math.fsum(baseline_national[k] *
                                 classification['score_by_old'][k][metric] for k in old) / denominator
                                for metric in ['broad', 'lm']}
        audit.append({'common_group': group, 'naics2017_members': ';'.join(old),
                      'naics2022_members': ';'.join(members['new']),
                      'aiie_common': fixed.get(group, {}).get('broad'),
                      'lm_aiie_common': fixed.get(group, {}).get('lm'),
                      'score_status': reason or 'scored',
                      'is_changed_group': int(members['changed'])})
    return fixed, audit


def sector_code(industry):
    prefix = industry[:2]
    return {'31': '31-33', '32': '31-33', '33': '31-33',
            '44': '44-45', '45': '44-45', '48': '48-49', '49': '48-49'}.get(prefix, prefix)


def fixed_sector_scores(classification, national):
    """较粗的稳健性代理；不利用各县未披露四位行业构成。"""
    members = defaultdict(list)
    for group in classification['group_members'].values():
        if len({sector_code(k) for k in group['old'] + group['new']}) != 1:
            raise ValueError('跨版本行业组跨越多个大类，不能计算固定行业大类代理。')
    for k in classification['old_to_group']:
        members[sector_code(k)].append(k)
    scores, audit = {}, []
    for sector, codes in sorted(members.items()):
        unscored = [k for k in codes if k not in classification['score_by_old']]
        positive = [k for k in codes if national.get(k) is not None and national[k] > 0]
        # 大类内有未评分产业即不赋予整个大类分数，避免补造农业/家庭就业暴露。
        reason = 'unscored_member_industry' if unscored else ''
        if not reason and any(national.get(k) is None for k in codes):
            reason = 'missing_national_weight'
        denominator = sum(national[k] for k in positive)
        if not reason and denominator > 0:
            scores[sector] = {metric: math.fsum(national[k] * classification['score_by_old'][k][metric]
                                             for k in positive) / denominator for metric in ['broad', 'lm']}
        elif not reason:
            reason = 'zero_national_weight'
        audit.append({'sector_code': sector, 'naics2017_members': ';'.join(sorted(codes)),
                      'aiie_sector': scores.get(sector, {}).get('broad'),
                      'lm_aiie_sector': scores.get(sector, {}).get('lm'),
                      'score_status': reason or 'scored'})
    return scores, audit


def sector_index(sectors, scores, total):
    matched = sum(e for k, e in sectors.items() if e is not None and k in scores)
    return {'ai_exposure_sector_proxy': math.fsum(e * scores[k]['broad'] for k, e in sectors.items()
            if e is not None and k in scores) / matched if matched > 0 else None,
            'lm_exposure_sector_proxy': math.fsum(e * scores[k]['lm'] for k, e in sectors.items()
            if e is not None and k in scores) / matched if matched > 0 else None,
            'sector_matched_employment': matched,
            'sector_employment_coverage': matched / total if total is not None and total > 0 else None,
            'suppressed_sector_count': sum(e is None for e in sectors.values())}


def county_index(industry_employment, mapping, scores, total):
    observed = sum(e for e in industry_employment.values() if e is not None)
    matched, broad, lm, unscored, unclassified = 0, [], [], 0, 0
    groups, suppressed, unmapped = set(), 0, set()
    for industry, e in industry_employment.items():
        # 9999是QCEW未分类行业，不是官方NAICS；保留其就业，纳入覆盖率缺口。
        if industry == '9999':
            if e is None:
                suppressed += 1
            else:
                unscored += e
                unclassified += e
            continue
        if industry not in mapping:
            unmapped.add(industry)
            continue
        group = mapping[industry]
        if e is None:
            suppressed += 1
            continue
        if group not in scores:
            unscored += e
            continue
        matched += e
        broad.append(e * scores[group]['broad'])
        lm.append(e * scores[group]['lm'])
        if e > 0:
            groups.add(group)
    if unmapped:
        raise ValueError(f'QCEW四位行业没有官方跨分类映射：{sorted(unmapped)}')
    coverage = matched / total if total is not None and total > 0 else None
    return {'ai_exposure': math.fsum(broad) / matched if matched > 0 else None,
            'lm_ai_exposure': math.fsum(lm) / matched if matched > 0 else None,
            'private_employment': total, 'matched_employment': matched,
            'disclosed_4digit_employment': observed, 'unscored_disclosed_employment': unscored,
            'unclassified_employment': unclassified,
            'employment_coverage': coverage, 'matched_group_count': len(groups),
            'suppressed_industry_count': suppressed,
            'coverage_over_101pct': int(coverage is not None and coverage > 1.01),
            'data_status': 'estimated_from_disclosed_industries' if matched > 0 else 'no_matched_employment'}


def make_panel(extractions, classification, scores, names, threshold, sector_scores=None):
    panel = []
    for data in extractions:
        year = data['year']
        mapping = classification['old_to_group'] if year == 2021 else classification['new_to_group']
        for fips, total in sorted(data['county_totals'].items()):
            if fips not in names:
                raise ValueError(f'地区名称缺失：{fips}')
            name = names[fips]
            combined = int('includes' in name.lower() or fips == '15009')
            unstable = int(fips.startswith('09'))
            row = {'activity_year': year, 'county_code': fips, 'state_fips': fips[:2],
                   'qcew_area_name': name, 'naics_version': 2017 if year == 2021 else 2022,
                   'ct_geography': ('planning_region' if year == 2024 else 'old_county') if unstable else '',
                   'geography_break': unstable, 'combined_area': combined}
            row.update(county_index(data['county_industries'].get(fips, {}), mapping, scores, total))
            row.update(sector_index(data.get('county_sectors', {}).get(fips, {}), sector_scores or {}, total))
            coverage = row['employment_coverage']
            row['coverage_ge_threshold'] = int(coverage is not None and coverage >= threshold)
            row['recommended_sample'] = int(row['coverage_ge_threshold'] and not unstable and not combined
                                            and not row['coverage_over_101pct'])
            sector_coverage = row['sector_employment_coverage']
            row['sector_recommended_sample'] = int(sector_coverage is not None and threshold <= sector_coverage <= 1.01
                                                   and not unstable and not combined)
            panel.append(row)
        # Kalawao并入Maui发布，保留空值，绝不伪造单独的县暴露度。
        if '15005' not in data['county_totals']:
            row = {'activity_year': year, 'county_code': '15005', 'state_fips': '15',
                   'qcew_area_name': 'Kalawao County, Hawaii (not published separately)',
                   'naics_version': 2017 if year == 2021 else 2022,
                   'ct_geography': '', 'geography_break': 0, 'combined_area': 1}
            row.update(county_index({}, mapping, scores, None))
            row.update(sector_index({}, sector_scores or {}, None))
            row.update(data_status='not_published_separately', coverage_ge_threshold=0, recommended_sample=0,
                       sector_recommended_sample=0)
            panel.append(row)
    panel.sort(key=lambda r: (r['county_code'], r['activity_year']))
    keys = [(r['county_code'], r['activity_year']) for r in panel]
    if len(set(keys)) != len(keys):
        raise ValueError('县×年主键重复。')
    # 使用一次性2021县横截面的均值/样本标准差，不逐年重新标准化。
    reference = [r for r in panel if r['activity_year'] == 2021 and r['recommended_sample']]
    parameters = {'reference_year': 2021, 'coverage_threshold': threshold, 'reference_counties': len(reference)}
    for metric in ['ai_exposure', 'lm_ai_exposure', 'ai_exposure_sector_proxy', 'lm_exposure_sector_proxy']:
        metric_reference = ([r for r in panel if r['activity_year'] == 2021 and r['sector_recommended_sample']]
                            if 'sector_proxy' in metric else reference)
        values = [r[metric] for r in metric_reference if r[metric] is not None]
        if not values and sector_scores is None:
            continue
        if len(values) < 2 or statistics.stdev(values) <= 0:
            raise ValueError('2021标准化参考样本不足。')
        mean, sd = statistics.mean(values), statistics.stdev(values)
        parameters[metric] = {'mean': mean, 'sample_sd': sd, 'reference_counties': len(values)}
        for r in panel:
            r[metric + '_z2021'] = (r[metric] - mean) / sd if r[metric] is not None else None
    return panel, parameters


def export_dta(panel, path):
    try:
        import pandas as pd
    except ImportError:
        LOGGER.info('未安装pandas；已输出CSV，可在Stata直接导入。')
        return False
    frame = pd.DataFrame(panel)
    text_columns = {'county_code', 'state_fips', 'qcew_area_name', 'ct_geography', 'data_status'}
    for column in frame:
        if column not in text_columns:
            frame[column] = pd.to_numeric(frame[column], errors='raise')
    temporary = path.with_name(path.name + '.part')
    frame.to_stata(temporary, write_index=False, version=118, variable_labels={
        'ai_exposure': 'Broad AI exposure, observed private employment weighted',
        'lm_ai_exposure': 'Language-model exposure, observed private employment weighted',
        'employment_coverage': 'Matched employment / county total private employment',
        'recommended_sample': 'Coverage threshold met, stable uncombined geography',
    })
    os.replace(temporary, path)
    return True


def main():
    parser = argparse.ArgumentParser(description='下载并计算2021—2024县级AI暴露度（私营就业口径）。')
    parser.add_argument('--output-dir', type=Path, default=Path(__file__).resolve().parent / 'AI_exposure_2021_2024')
    parser.add_argument('--coverage-threshold', type=float, default=0.90,
                        help='质量筛选参考值，不是数据准确性的保证；默认0.90。')
    parser.add_argument('--refresh', action='store_true', help='重新下载原始文件并重建就业缓存。')
    parser.add_argument('--offline', action='store_true', help='只使用已有且校验通过的原始文件。')
    parser.add_argument('--timeout', type=float, default=120)
    parser.add_argument('--retries', type=int, default=3)
    args = parser.parse_args()
    if not 0 < args.coverage_threshold <= 1 or args.retries < 1 or args.timeout <= 0:
        parser.error('阈值须在(0,1]；timeout/retries须为正数。')
    if args.offline and args.refresh:
        parser.error('--offline不能与--refresh同时使用。')
    out = args.output_dir.resolve()
    raw, processed = out / 'raw', out / 'processed'
    raw.mkdir(parents=True, exist_ok=True)
    processed.mkdir(parents=True, exist_ok=True)
    receipts = {}
    for name, url in SOURCES.items():
        target = raw / name
        if args.offline and not target.exists():
            raise ValueError(f'离线模式缺少原始文件：{target}')
        receipts[name] = fetch(url, target, args.timeout, args.retries, args.refresh)
    classification = load_classification(raw)
    names = read_area_titles(raw / 'area_titles.csv')
    extractions = [extract_qcew(raw / f'{y}_annual_singlefile.zip', y, processed,
                   receipts[f'{y}_annual_singlefile.zip']['sha256'], args.refresh) for y in range(2021, 2025)]
    scores, group_audit = fixed_group_scores(classification, extractions[0]['national'])
    sector_scores, sector_audit = fixed_sector_scores(classification, extractions[0]['national'])
    panel, standardization = make_panel(extractions, classification, scores, names, args.coverage_threshold, sector_scores)
    write_csv(out / 'county_ai_exposure_2021_2024.csv', list(panel[0]), panel)
    dta_written = export_dta(panel, out / 'county_ai_exposure_2021_2024.dta')
    write_csv(out / 'common_industry_scores.csv', list(group_audit[0]), group_audit)
    write_csv(out / 'sector_proxy_scores.csv', list(sector_audit[0]), sector_audit)
    write_json(out / 'classification_audit.json', classification)
    write_csv(out / 'published_score_mapping.csv', list(classification['score_rows'][0]), classification['score_rows'])
    summary = []
    for year in range(2021, 2025):
        rows = [r for r in panel if r['activity_year'] == year]
        c = [r['employment_coverage'] for r in rows if r['employment_coverage'] is not None]
        summary.append({'activity_year': year, 'county_rows': len(rows),
                        'valid_exposure_rows': sum(r['ai_exposure'] is not None for r in rows),
                        'coverage_threshold_rows': sum(r['coverage_ge_threshold'] for r in rows),
                        'recommended_sample_rows': sum(r['recommended_sample'] for r in rows),
                        'valid_sector_proxy_rows': sum(r['ai_exposure_sector_proxy'] is not None for r in rows),
                        'sector_recommended_rows': sum(r['sector_recommended_sample'] for r in rows),
                        'median_sector_coverage': statistics.median([r['sector_employment_coverage'] for r in rows
                                                                   if r['sector_employment_coverage'] is not None]),
                        'median_employment_coverage': statistics.median(c),
                        'min_employment_coverage': min(c), 'max_employment_coverage': max(c),
                        'suppressed_industry_rows': sum(r['suppressed_industry_count'] for r in rows)})
    write_csv(out / 'annual_summary.csv', list(summary[0]), summary)
    write_json(out / 'run_metadata.json', {'created_at_utc': utc_now(), 'years': [2021, 2022, 2023, 2024],
        'scope': '50 states and DC; QCEW workplace-based private covered employment; ownership 5',
        'formula': 'sum(disclosed matched industry employment * fixed common-group score) / sum(disclosed matched employment)',
        'coverage_formula': 'matched employment / total county private employment',
        'harmonization': 'Connected components of official NAICS2017-to-2022 four-digit concordance; fixed 2021 national private employment weights within groups',
        'scores_time_invariant': True, 'standardization': standardization,
        'sector_proxy_method': 'Coarse robustness proxy: fixed 2021 national private 4-digit weighted sector score, weighted by disclosed county sector employment; exclude any sector with an unscored member (11,81,92), and unclassified 99',
        'qcew_ct_switch_year': 2024, 'kalawao_not_separate': True,
        'unscored_groups': sum(r['score_status'] != 'scored' for r in group_audit),
        'sources': receipts, 'annual_summary': summary, 'stata_written': dta_written,
        'files': {name: sha256_file(out / name) for name in [
            'county_ai_exposure_2021_2024.csv', 'common_industry_scores.csv',
            'published_score_mapping.csv', 'sector_proxy_scores.csv', 'annual_summary.csv']}})
    LOGGER.info('完成：%s条县×年记录。结果保存于：%s', len(panel), out)
    for row in summary:
        LOGGER.info('%s年：%s条记录，%s条有暴露度，覆盖率达阈值%s条。', row['activity_year'],
                    row['county_rows'], row['valid_exposure_rows'], row['coverage_threshold_rows'])


if __name__ == '__main__':
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    logging.basicConfig(level=logging.INFO, format='%(asctime)s %(message)s')
    try:
        main()
    except (Exception, KeyboardInterrupt) as exc:
        LOGGER.error('%s: %s', type(exc).__name__, exc)
        sys.exit(1)

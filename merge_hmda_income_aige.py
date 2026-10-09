from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import pandas as pd


def normalize_fips(value: object) -> object:
    if pd.isna(value):
        return pd.NA

    text = str(value).strip()
    if not text:
        return pd.NA
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text.zfill(5)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Keep HMDA observations for 2021-2024 and merge county income "
            "and Appendix C AIGE values by county FIPS."
        )
    )
    parser.add_argument(
        "--dta",
        type=Path,
        default=Path("hmda_2018_2025_applicant_age_under_25.dta"),
        help="Input HMDA Stata file.",
    )
    parser.add_argument(
        "--income",
        type=Path,
        default=Path("county_per_capita_income_acs5_US_2021_2022_2023_2024.csv"),
        help="Input county income CSV file.",
    )
    parser.add_argument(
        "--appendix",
        type=Path,
        default=Path("AIOE_DataAppendix.xlsx"),
        help="AIOE workbook containing the Appendix C sheet.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("hmda_2021_2024_income_aige_merged.dta"),
        help="Output Stata file. An existing file will be replaced.",
    )
    parser.add_argument(
        "--chunksize",
        type=int,
        default=200_000,
        help="Number of input rows read per chunk (default: 200000).",
    )
    args = parser.parse_args()

    for path in (args.dta, args.income, args.appendix):
        if not path.is_file():
            raise FileNotFoundError(f"Input file not found: {path}")
    if args.chunksize < 1:
        raise ValueError("--chunksize must be a positive integer")
    if args.output.resolve() in {
        args.dta.resolve(),
        args.income.resolve(),
        args.appendix.resolve(),
    }:
        raise ValueError("The output path must not overwrite an input file.")

    income = pd.read_csv(args.income, dtype={"county_fips": str})
    income = income[
        ["year", "county_fips", "per_capita_income_usd", "moe_usd"]
    ].copy()
    income["county_fips"] = income["county_fips"].map(normalize_fips)
    income = income.rename(
        columns={
            "year": "activity_year",
            "moe_usd": "per_capita_income_moe_usd",
        }
    )

    workbook = openpyxl.load_workbook(
        args.appendix, read_only=True, data_only=True
    )
    try:
        if "Appendix C" not in workbook.sheetnames:
            raise ValueError(f"{args.appendix} does not contain an Appendix C sheet")
        worksheet = workbook["Appendix C"]
        aige_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()

    aige = pd.DataFrame(
        aige_rows, columns=["FIPS Code", "Geographic Area", "AIGE"]
    )[["FIPS Code", "AIGE"]].copy()
    aige["FIPS Code"] = aige["FIPS Code"].map(normalize_fips)
    aige = aige.rename(columns={"FIPS Code": "county_fips"})

    reader = pd.read_stata(
        args.dta,
        iterator=True,
        chunksize=args.chunksize,
        convert_categoricals=False,
    )
    parts = []
    source_rows = 0
    selected_rows = 0

    for chunk in reader:
        source_rows += len(chunk)
        chunk = chunk.loc[chunk["activity_year"].between(2021, 2024)].copy()
        selected_rows += len(chunk)
        if chunk.empty:
            continue

        chunk["_county_fips"] = chunk["county_code"].astype(str).str.strip()
        chunk.loc[chunk["_county_fips"].eq(""), "_county_fips"] = pd.NA
        chunk["_county_fips"] = chunk["_county_fips"].map(normalize_fips)

        chunk = chunk.merge(
            income,
            left_on=["activity_year", "_county_fips"],
            right_on=["activity_year", "county_fips"],
            how="left",
            validate="many_to_one",
            sort=False,
        ).drop(columns=["county_fips"])
        chunk = chunk.merge(
            aige,
            left_on="_county_fips",
            right_on="county_fips",
            how="left",
            validate="many_to_one",
            sort=False,
        ).drop(columns=["_county_fips", "county_fips"])
        parts.append(chunk)

    if selected_rows == 0:
        raise ValueError("No HMDA observations found for activity_year 2021-2024.")

    merged = pd.concat(parts, ignore_index=True)
    if len(merged) != selected_rows:
        raise RuntimeError(
            f"Row count changed during merge: {len(merged)} != {selected_rows}"
        )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary_output = args.output.with_name(
        f"{args.output.stem}.tmp{args.output.suffix}"
    )
    try:
        merged.to_stata(temporary_output, write_index=False, version=118)
        temporary_output.replace(args.output)
    finally:
        if temporary_output.exists():
            temporary_output.unlink()

    print(f"Output: {args.output.resolve()}")
    print(f"Source rows: {source_rows}")
    print(f"Rows retained (activity_year 2021-2024): {selected_rows}")
    print(f"Income matched: {merged['per_capita_income_usd'].notna().sum()}")
    print(f"AIGE matched: {merged['AIGE'].notna().sum()}")
    print(f"Years: {sorted(merged['activity_year'].unique().tolist())}")


if __name__ == "__main__":
    main()
